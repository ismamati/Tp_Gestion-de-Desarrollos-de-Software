"""Genera diapositivas editables y un guion a partir del mismo contenido.

Herramienta de documentación: requiere python-pptx, no es parte de la aplicación.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


SLIDES = [
    {
        "title": "Que funcione es el comienzo",
        "subtitle": "Calidad de código, deuda técnica y seguridad · DevSecOps",
        "left_title": "Nuestra propuesta",
        "left": ["Comprobar comportamiento", "Revisar calidad y seguridad", "Controlar antes de integrar"],
        "right_title": "Una demo pequeña",
        "right": ["Notas de alumnos en Python", "GitHub Actions + SonarCloud", "Dependabot + protección de main"],
        "notes": "Completar los nombres del equipo. Presentamos el jueves 15 de octubre de 2026. La idea central: un cambio puede pasar los tests y aun así necesitar mejoras antes de integrarse. No afirmamos que los controles garanticen ausencia total de defectos.",
    },
    {
        "title": "¿Qué hace la aplicación?",
        "subtitle": "Un módulo Python; los datos entran como una lista de alumnos",
        "left_title": "Tres funciones",
        "left": ["Buscar alumnos por materia", "Aprobar con nota ≥ 6", "Resumir cantidad, promedio y mejor nota"],
        "right_title": "Un ejemplo concreto",
        "right": ["Ana: 8 · Bruno: 4 · Carla: 9", "Aprobados: Ana y Carla", "Promedio del curso: 7,0"],
        "notes": "No es una web: no tiene pantallas, servidor ni base de datos. app/notas.py recibe datos y devuelve resultados. buscar_por_materia filtra; esta_aprobado devuelve verdadero o falso; reporte_curso devuelve un diccionario. Los cinco tests están en tests/test_notas.py. pytest corre los tests y pytest-cov mide cobertura.",
    },
    {
        "title": "La deuda aparece al mantener",
        "subtitle": "Una solución rápida puede aumentar el costo del próximo cambio",
        "left_title": "Deuda técnica",
        "left": ["Copiamos el cálculo de un reporte", "Una corrección requiere cambiar dos sitios", "Más trabajo y riesgo de inconsistencias"],
        "right_title": "Code Smells",
        "right": ["Variables que nunca se usan", "Funciones difíciles de entender", "Señales para revisar la mantenibilidad"],
        "notes": "Un Code Smell no significa que el programa ya falla ni que hay una vulnerabilidad. SonarCloud mide la duplicación por separado. La deuda puede ser consciente si se registra y se planifica su corrección. En nuestro PR hay una copia de reporte_curso y una variable total_inscriptos que no se utiliza.",
    },
    {
        "title": "Dos formas de revisar el código",
        "subtitle": "Los tests y el análisis estático responden preguntas diferentes",
        "left_title": "Testing dinámico",
        "left": ["Ejecuta el código", "Compara resultados con lo esperado", "pytest aquí; Playwright en navegador"],
        "right_title": "Análisis estático · SAST",
        "right": ["Inspecciona sin ejecutar la aplicación", "Busca patrones de errores y seguridad", "SonarCloud analiza nuestro cambio"],
        "notes": "Ejemplo: pytest llama a esta_aprobado y comprueba el resultado. SonarCloud inspecciona el código y señala patrones problemáticos. SonarCloud no ejecuta nuestros tests: Actions ejecuta pytest y SonarCloud importa coverage.xml. Ninguno de estos controles detecta todos los problemas.",
    },
    {
        "title": "También dependemos de código ajeno",
        "subtitle": "La cadena de suministro incluye librerías y herramientas del CI",
        "left_title": "Qué puede pasar",
        "left": ["Una versión tiene una vulnerabilidad", "Una cuenta de publicación es comprometida", "Un paquete malicioso llega al proyecto"],
        "right_title": "Dependabot",
        "right": ["Detecta versiones con avisos conocidos", "Genera alertas y propone actualizaciones", "Ya tenemos el PR #1 de python-multipart"],
        "notes": "Una vulnerabilidad conocida no prueba que haya existido un ataque. requirements.txt incluye python-multipart 0.0.6, afectado por CVE-2024-24762. El módulo de notas no lo usa; está como ejemplo. Dependabot ya propuso actualizar a 0.0.31. Una alerta de Dependabot no hace fallar automáticamente el gate de SonarCloud. Fuente: https://github.com/advisories/GHSA-2jv5-9r88-3w3p .",
    },
    {
        "title": "Qué tenemos y qué falta",
        "subtitle": "Estado consultado el 13/10/2026 · comprobar antes de grabar",
        "left_title": "Ya funciona",
        "left": ["Tests y reportes de cobertura", "SonarCloud dentro del workflow", "PR #2 rojo y PR #1 de Dependabot"],
        "right_title": "Por completar",
        "right": ["Verificar bloqueo real de merge", "Corregir y comprobar checks verdes", "Grabar evidencia y ensayar"],
        "notes": "La protección de main no pudo verificarse con la API pública porque requiere autenticación. No presentarla como configurada hasta comprobarla. El PR #2 está abierto y falla por el Quality Gate; el PR #1 también está abierto. La corrección verde todavía no está hecha en esta preparación.",
    },
    {
        "title": "Qué recorre cada Pull Request",
        "subtitle": "GitHub Actions ejecuta la secuencia automáticamente",
        "flow": ["PR", "Instalar", "Tests +\ncobertura", "SonarCloud", "Gate", "Merge*"],
        "left_title": "Quality Gate",
        "left": ["Bugs nuevos = 0", "Vulnerabilidades nuevas = 0", "Cobertura nueva ≥ 80%"],
        "right_title": "*Checks obligatorios",
        "right": ["GitHub debe exigir su aprobación", "Gate rojo → check rojo → bloqueo", "Otras condiciones también se cumplen"],
        "notes": "El workflow genera coverage.xml después de ejecutar pytest. El step oficial de SonarCloud ejecuta SonarScanner y espera el gate en PR con sonar.qualitygate.wait=true. En pushes a main el workflow actual no espera el resultado del gate. La regla de duplicación observada tiene un máximo de 3%. En un PR se evalúa código nuevo. Para cambios menores de 20 líneas, desmarcar Ignore duplication and coverage on small changes en Quality Gate si se quiere exigir ese control. Fuente: https://docs.sonarsource.com/sonarqube-cloud/managing-your-projects/project-analysis/changing-quality-gate .",
    },
    {
        "title": "Leer el dashboard",
        "subtitle": "Comparar el código global con el código nuevo del PR",
        "left_title": "Problemas detectados",
        "left": ["Bugs → confiabilidad", "Vulnerabilidades → seguridad", "Code Smells → mantenibilidad"],
        "right_title": "Otras métricas",
        "right": ["Duplicación → bloques repetidos", "Cobertura → código ejecutado en tests", "100% no prueba todos los escenarios"],
        "notes": "Mostrar el dashboard real y la configuración del gate. Para main se consultaron cobertura global de 100%, cero Bugs, Vulnerabilidades y Code Smells, y duplicación de 0%. No mezclar estas métricas con las del PR. Abrir https://sonarcloud.io/dashboard?id=ismamati_Tp_Gestion-de-Desarrollos-de-Software .",
    },
    {
        "title": "La demo roja ya está preparada",
        "subtitle": "PR #2 · Agrego reporte de aprobados",
        "left_title": "Qué ocurre",
        "left": ["Se copia el cálculo del reporte", "Los tests antiguos siguen pasando", "La función nueva casi no se ejecuta"],
        "right_title": "Resultado comprobado",
        "right": ["Cobertura nueva: 4,8% < 80%", "Duplicación: 82,8% > 3%", "Quality Gate: ERROR"],
        "capture": "Agregar captura real: gate rojo + checks + merge bloqueado",
        "notes": "Abrir https://github.com/ismamati/Tp_Gestion-de-Desarrollos-de-Software/pull/2 . Se detectó también un Code Smell, pero la condición de mantenibilidad pasa: no atribuir el fallo a la variable sin usar. Los motivos comprobados son cobertura y duplicación. Mostrar el diff y luego los checks. Registrar la evidencia roja antes de modificar la rama. El bloqueo de merge necesita checks obligatorios.",
    },
    {
        "title": "La corrección cabe en tres líneas",
        "subtitle": "Filtrar los aprobados y reutilizar el reporte existente",
        "code": "def reporte_aprobados(alumnos):\n    aprobados = [a for a in alumnos if esta_aprobado(a)]\n    return reporte_curso(aprobados)",
        "left_title": "Qué mejora",
        "left": ["Un solo cálculo para mantener", "Sin la variable que no se usa", "Mismo objetivo funcional"],
        "right_title": "Qué pruebo",
        "right": ["Ana y Carla: promedio 8,5", "Sin aprobados: reporte vacío", "Una nota de 6 también aprueba"],
        "notes": "Usar la misma rama demo-code-smell del PR #2. Reemplazar toda la función copiada por estas líneas y agregar tests a tests/test_notas.py. Ejecutar pytest con cobertura, hacer commit y push. El código usa a como nombre corto en pantalla; la guía usa alumno para mayor claridad. No editar main para corregir un PR.",
    },
    {
        "title": "Cerrar el ciclo con evidencia",
        "subtitle": "El resultado verde debe comprobarse después del nuevo push",
        "left_title": "En GitHub y SonarCloud",
        "left": ["Mismo PR, nuevo commit", "Tests y gate aprobados", "Merge permitido al cumplir requisitos"],
        "right_title": "En Dependabot",
        "right": ["Mostrar la alerta de seguridad", "Mostrar la versión propuesta en PR #1", "Revisar checks antes de integrar"],
        "capture": "Agregar captura real del nuevo análisis verde y del PR de Dependabot",
        "notes": "Esta diapositiva describe el resultado pendiente, no uno que ya se haya producido. Verificar tests y SonarCloud Code Analysis verdes. Si siguen pendientes revisiones, conversaciones u otras condiciones, el merge puede continuar bloqueado. No usar bypass. Configurar la protección para aplicar a administradores si presentan desde esa cuenta. Dependabot: Settings → Advanced Security → Dependency graph, alerts y security updates. Fuente: https://docs.github.com/en/code-security/how-tos/secure-your-supply-chain/secure-your-dependencies/configure-security-updates .",
    },
    {
        "title": "Cada control responde una pregunta",
        "subtitle": "Automatizar la revisión para integrar mejores cambios",
        "left_title": "Antes de integrar",
        "left": ["Tests: ¿se comporta como esperamos?", "SonarCloud: ¿cumple calidad y seguridad?", "Dependabot: ¿hay versiones vulnerables?"],
        "right_title": "Lo que demostramos",
        "right": ["El rojo tiene un motivo concreto", "La corrección mejora código y tests", "GitHub exige los checks configurados"],
        "notes": "Cerrar con la idea de que calidad y seguridad forman parte del flujo habitual de desarrollo. Preparar respuestas: cobertura no equivale a ausencia de Bugs; un smell no es una vulnerabilidad; SonarCloud y Dependabot tienen funciones diferentes; un check rojo solo bloquea si es obligatorio. El archivo GUIA_EQUIPO.md tiene el procedimiento detallado y una alternativa pequeña llamada contar_aprobados.",
    },
]

NAVY = "142437"
TEAL = "007E87"
INK = "203449"
MUTED = "526678"
WHITE = "FFFFFF"
BG = "EFF4F8"


def box(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def label(slide, x, y, w, h, value, size=22, color=INK, bold=False, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(10)
    return shape


def main():
    out_dir = Path(__file__).resolve().parent
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Calidad de código, deuda técnica y seguridad"
    prs.core_properties.subject = "Introducción y demostración DevSecOps"
    prs.core_properties.language = "es-AR"
    markdown = ["# Diapositivas y notas para exponer", "", "Presentación: jueves 15/10/2026. Completar nombres y capturas reales.", ""]

    for number, data in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        box(slide, 0, 0, 13.333, 1.8, NAVY)
        box(slide, 0.48, 0.37, 0.08, 0.8, TEAL)
        label(slide, 0.78, 0.32, 11.85, 0.7, data["title"], 34, WHITE, True)
        label(slide, 0.8, 1.12, 11.8, 0.5, data["subtitle"], 18, "C9DAE7")

        top = 2.25
        if "flow" in data:
            for index, word in enumerate(data["flow"]):
                x = 0.72 + index * 2.03
                box(slide, x, 2.13, 1.72, 0.9, TEAL)
                label(slide, x + 0.13, 2.24, 1.47, 0.68, word, 18, WHITE, True)
                if index < 5:
                    label(slide, x + 1.77, 2.36, 0.25, 0.4, "›", 23, TEAL, True)
            top = 3.47

        if "code" in data:
            box(slide, 0.72, 2.14, 11.89, 1.42, NAVY)
            label(slide, 0.99, 2.34, 11.4, 1.0, data["code"], 19, "DBF4E9", font="Consolas")
            top = 3.91

        card_height = 3.55 if top < 3 else 2.54
        for side, x in [("left", 0.72), ("right", 6.81)]:
            box(slide, x, top, 5.8, card_height, WHITE)
            label(slide, x + 0.28, top + 0.23, 5.22, 0.48, data[side + "_title"], 24, TEAL, True)
            step = 0.77 if top < 3 else 0.57
            for index, item in enumerate(data[side]):
                label(slide, x + 0.28, top + 0.94 + index * step, 5.15, 0.68 if top < 3 else 0.53, "• " + item, 22 if top < 3 else 20)

        if "capture" in data:
            box(slide, 0.72, 6.06, 11.89, 0.5, "DCEAEF")
            label(slide, 0.93, 6.15, 11.4, 0.35, data["capture"], 15, MUTED)

        label(slide, 0.74, 7.0, 10.5, 0.25, "GDS  ·  DEVSECOPS  ·  15 OCT 2026", 11, MUTED)
        label(slide, 12.1, 6.96, 0.8, 0.3, f"{number:02d}", 13, TEAL, True)
        slide.notes_slide.notes_text_frame.text = data["notes"]

        markdown += [f"## {number}. {data['title']}", "", data["subtitle"], ""]
        if "flow" in data:
            markdown += [" → ".join(w.replace("\n", " ") for w in data["flow"]), ""]
        if "code" in data:
            markdown += ["```python", data["code"], "```", ""]
        for side in ["left", "right"]:
            markdown += [f"**{data[side + '_title']}**", ""] + ["- " + item for item in data[side]] + [""]
        if "capture" in data:
            markdown += ["**Captura pendiente:** " + data["capture"], ""]
        markdown += ["**Para explicar oralmente:** " + data["notes"], ""]

    target = out_dir / "PRESENTACION_DEVSECOPS.pptx"
    prs.save(target)
    (out_dir / "DIAPOSITIVAS.md").write_text("\n".join(markdown), encoding="utf-8")
    restored = Presentation(target)
    assert len(restored.slides) == len(SLIDES)
    for slide, expected in zip(restored.slides, SLIDES):
        assert expected["title"] in "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert slide.notes_slide.notes_text_frame.text == expected["notes"]
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= restored.slide_width
            assert shape.top + shape.height <= restored.slide_height
    print(f"Generadas y verificadas {len(restored.slides)} diapositivas con notas: {target}")


if __name__ == "__main__":
    main()
